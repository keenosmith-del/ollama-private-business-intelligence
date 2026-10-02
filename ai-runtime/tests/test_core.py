import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from app import main
from app.main import chunk_text, classify, extract_text, parse_sales_csv, vector_literal
def test_chunk_text_overlaps_and_preserves_content():
    chunks=chunk_text('word '*600,size=300,overlap=40)
    assert len(chunks)>1 and all(chunks)
def test_classifier_routes_hybrid_question():
    assert classify('Why did revenue fall in Q2?')=='hybrid'
    assert classify('Which customers contributed most to the revenue decline and what does management say?')=='hybrid'
def test_classifier_separates_structured_financial_queries():
    assert classify('What are our biggest areas of financial loss?')=='analytics'
    assert classify('Which customers declined the most in Q2?')=='analytics'
    assert classify('Which customers are becoming less profitable?')=='analytics'
def test_vector_literal_enforces_database_dimension():
    assert vector_literal([0.0]*768).startswith('[0.0,')
    with pytest.raises(ValueError): vector_literal([0.0]*3)
def test_plain_text_extraction_and_unsupported_file_rejection():
    assert extract_text('policy.txt',b'Refunds are accepted within 30 days.')==[(None,'Refunds are accepted within 30 days.')]
    with pytest.raises(HTTPException) as error: extract_text('file.exe',b'no')
    assert error.value.status_code==415
def test_sales_csv_parser_validates_and_normalizes_rows():
    rows=parse_sales_csv(b'customer,date,amount,industry\nAcme Retail,2026-09-01,1250.50,Retail\n')
    assert rows[0]['customer']=='Acme Retail' and str(rows[0]['amount'])=='1250.50'
    with pytest.raises(HTTPException) as error: parse_sales_csv(b'customer,date,amount\nAcme,not-a-date,1\n')
    assert error.value.status_code==422
    with pytest.raises(HTTPException): parse_sales_csv(b'customer,date,amount\nAcme,2026-09-01,0.001\n')
def test_ai_routes_require_gateway_token(monkeypatch):
    monkeypatch.setattr(main,'INTERNAL_TOKEN','a-test-internal-token-with-more-than-32-chars')
    response=TestClient(main.app).post('/v1/analysis',headers={'x-org-id':'org-test'},json={'question':'Why did revenue fall?'})
    assert response.status_code==401

def test_policy_is_knowledge_even_when_it_mentions_expenses():
    assert classify('What does our internal expense policy say?') == 'knowledge'


def test_period_selection_uses_requested_quarters_and_year():
    from datetime import date
    from app.analytics import select_periods
    previous, target, quarters, _ = select_periods('Compare Q1 and Q2 revenue', date(2027, 9, 1))
    assert previous == date(2027, 1, 1) and target == date(2027, 4, 1)
    assert quarters == [1, 2]
    previous, target, _, _ = select_periods('Q1 2028 quarter-over-quarter revenue', date(2027, 9, 1))
    assert previous == date(2027, 10, 1) and target == date(2028, 1, 1)


def test_csv_document_handles_extra_columns():
    assert 'column_3: extra' in extract_text('rows.csv', b'a,b\n1,2,extra\n')[0][1]


def test_expense_validation_rejects_nonfinite_and_missing_categories():
    from app.main import parse_expense_csv
    assert parse_expense_csv(b'category,date,amount\nDelivery,2027-04-01,15.25\n')[0]['amount'] == __import__('decimal').Decimal('15.25')
    for raw in (b'category,date,amount\nDelivery,2027-04-01,NaN\n', b'category,date,amount\n,2027-04-01,12\n'):
        with pytest.raises(HTTPException): parse_expense_csv(raw)


def test_office_formats_extract_table_and_sheet_content():
    import io
    from docx import Document
    from openpyxl import Workbook
    doc=Document();doc.add_paragraph('Synthetic audit policy')
    doc.add_table(rows=1, cols=1).cell(0,0).text='Receipts required above 500'
    buffer=io.BytesIO();doc.save(buffer)
    assert 'Receipts required above 500' in extract_text('policy.docx',buffer.getvalue())[0][1]
    book=Workbook();book.active.append(['category','amount']);book.active.append(['Delivery',25])
    buffer=io.BytesIO();book.save(buffer)
    assert 'amount: 25' in extract_text('costs.xlsx',buffer.getvalue())[0][1]

@pytest.mark.anyio
async def test_model_prose_and_invented_evidence_ids_are_not_trusted(monkeypatch):
    import httpx
    async def fake_post(self,*args,**kwargs):
        return httpx.Response(200,json={'message':{'content':'Revenue declined due to policy noncompliance costing 285000.'}},request=httpx.Request('POST','http://localhost/api/chat'))
    monkeypatch.setattr(httpx.AsyncClient,'post',fake_post)
    with pytest.raises((ValueError,KeyError)):
        await main.generate('Only select evidence IDs.')

@pytest.fixture
def anyio_backend():
    return 'asyncio'
